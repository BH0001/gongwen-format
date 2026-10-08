"""Focused regressions for numbering, content guards, preserved styles and mapping."""
import argparse
import json
import subprocess
import sys
import zipfile
from pathlib import Path
from docx import Document
from lxml import etree as ET
from test_gongwen import fixture

ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT / "scripts/gongwen.py"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def el(name, **attrs):
    n = ET.Element("{" + W + "}" + name)
    for k, v in attrs.items():
        n.set("{" + W + "}" + k, str(v))
    return n


def read(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def tree(files, part="word/document.xml"):
    return ET.fromstring(files[part])


def save(path, files, root=None):
    if root is not None:
        files["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True, standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)


def cmd(*args, fail=False):
    r = subprocess.run([sys.executable, str(CLI), *map(str, args)], capture_output=True,
                       text=True, encoding="utf-8")
    assert (r.returncode != 0) if fail else (r.returncode == 0), r.stdout + r.stderr
    return r.stderr if fail else json.loads(r.stdout)


def mapping(work, source, index, **changes):
    info = cmd("inspect", source)
    p = info["paragraphs"][index]
    result = {"version": 1, "sourceSha256": info["sourceSha256"], "paragraphs": [
        {"id": p["id"], "expectedText": p["expectedText"], **changes}]}
    path = work / (source.stem + "-map-" + str(len(list(work.glob("*-map-*.json")))) + ".json")
    path.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--work", required=True)
    work = Path(ap.parse_args().work).resolve()
    work.mkdir(parents=True, exist_ok=False)
    passed = []
    source = work / "numbering.docx"
    doc = Document(); doc.add_paragraph("编号检查通知", style="Title")
    doc.add_paragraph("编号必须保留。", style="List Number"); doc.save(source)
    files = read(source); root = tree(files)
    pp = root.findall("w:body/w:p", NS)[1].find("w:pPr", NS)
    np = el("numPr"); np.append(el("ilvl", val=0)); pp.append(np); save(source, files, root)
    out = work / "numbering-out.docx"; cmd("format", source, "--output", out)
    p = tree(read(out)).findall("w:body/w:p", NS)[1]
    assert p.find("w:pPr/w:numPr/w:numId", NS) is not None
    cmd("check", out, "--against", source)
    files = read(out); root = tree(files); np = root.findall("w:body/w:p", NS)[1].find("w:pPr/w:numPr", NS)
    np.remove(np.find("w:numId", NS)); bad = work / "lost-numbering.docx"; save(bad, files, root)
    cmd("check", bad, "--against", source, fail=True)
    passed.append("partial inherited numbering is resolved and numbering loss is detected")

    source = work / "controls.docx"
    doc = Document(); doc.add_paragraph("范围检查通知", style="Title")
    p = doc.add_paragraph("\t（一）𠮷字题头。正文。\n末行。")
    p._p.find("w:r", NS).insert(0, el("lastRenderedPageBreak")); doc.save(source)
    info = cmd("inspect", source)
    assert info["paragraphs"][1]["runs"][0]["start"] == 1
    out = work / "controls-out.docx"; cmd("format", source, "--output", out)
    root = tree(read(out))
    assert "".join(root.xpath("//w:r[w:rPr/w:rStyle[@w:val='GW_char_heading2']]/w:t/text()", namespaces=NS)) == "（一）𠮷字题头。"
    cmd("check", out, "--against", source)
    for name, xpath in [("tab", "//w:tab"), ("break", "//w:br")]:
        files = read(out); root = tree(files)
        for node in root.xpath(xpath, namespaces=NS): node.getparent().remove(node)
        bad = work / ("removed-" + name + ".docx"); save(bad, files, root)
        cmd("check", bad, "--against", source, fail=True)
    text = info["paragraphs"][1]["expectedText"]
    mp = mapping(work, source, 1, role="body", spans=[dict(start=0, length=len(text), role="heading2", expectedText=text)])
    cmd("format", source, "--output", work / "invalid-range.docx", "--map", mp, fail=True)
    passed.append("mixed runs retain Unicode offsets and controls; spans crossing controls are refused")

    # A generated title break is the only extra break allowed by comparison.
    source = work / "title.docx"
    doc = Document(); doc.add_paragraph("甲𠮷乙丙丁戊", style="Title"); doc.save(source)
    mp = mapping(work, source, 0, role="title", lineBreaks=[2])
    out = work / "title-out.docx"; cmd("format", source, "--output", out, "--map", mp)
    cmd("check", out, "--against", source)
    assert tree(read(out)).xpath("//w:r[w:rPr/w:rStyle[@w:val='GW_LayoutBreak']]/w:br", namespaces=NS)
    again = work / "title-twice.docx"; cmd("format", out, "--output", again)
    cmd("check", again, "--against", out)
    mp = mapping(work, source, 0, role="title", lineBreaks=[2], spans=[dict(start=3, length=2, role="heading2", expectedText="丙丁")])
    cmd("format", source, "--output", work / "invalid-title-span.docx", "--map", mp, fail=True)
    passed.append("title breaks are marked; combining original coordinates with new breaks is safely refused")

    source = work / "field.docx"
    doc = Document(); doc.add_paragraph("字段检查通知", style="Title"); p = doc.add_paragraph("来源：")
    field = el("fldSimple", instr=" DATE "); run = el("r"); t = el("t"); t.text = "2026年10月8日"
    run.append(t); field.append(run); p._p.append(field); doc.save(source)
    out = work / "field-out.docx"; cmd("format", source, "--output", out)
    files = read(out); root = tree(files)
    for field in root.xpath("//w:fldSimple", namespaces=NS):
        for child in list(field): field.addprevious(child)
        field.getparent().remove(field)
    bad = work / "removed-field.docx"; save(bad, files, root)
    cmd("check", bad, "--against", source, fail=True)
    passed.append("field structure cannot be replaced by identical plain text")

    source = work / "rich.docx"; fixture(source)
    out = work / "rich-out.docx"; cmd("format", source, "--output", out)
    cmd("check", out, "--against", source)
    files = read(out); root = tree(files); root.find(".//w:ins", NS).set("{" + W + "}author", "changed")
    bad = work / "revision-author.docx"; save(bad, files, root)
    cmd("check", bad, "--against", source, fail=True)
    passed.append("rich objects survive formatting and revision identity changes are detected")

    # Changing the current profile must not overwrite styles needed by a preserved paragraph.
    mp = mapping(work, out, 3, role="preserve")
    preserved = work / "preserved-out.docx"
    cmd("format", out, "--output", preserved, "--map", mp, "--font", "body=宋体")
    cmd("check", preserved, "--against", out)
    before = tree(read(out)).findall("w:body/w:p", NS)[3]
    after = tree(read(preserved)).findall("w:body/w:p", NS)[3]
    assert after.find("w:pPr/w:pStyle", NS).get("{" + W + "}val").startswith("GWP_")
    assert before.xpath(".//w:rFonts/@w:eastAsia", namespaces=NS) == after.xpath(".//w:rFonts/@w:eastAsia", namespaces=NS)
    passed.append("preserved GW paragraphs retain old effective fonts across profile changes")

    plain = cmd("check", source)
    assert plain["formatStatus"] == "not-applicable" and plain["formatNotice"]
    profile = json.loads((ROOT / "assets/profiles/gbt9704.json").read_text(encoding="utf-8"))
    for key, value in [("lineMode", "unknown"), ("signatureMode", "unknown"), ("rows", 0), ("fixedLinePt", -1), ("marginLeftMm", 999)]:
        config = dict(profile); config[key] = value
        path = work / ("bad-" + key + ".json"); path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
        cmd("format", source, "--output", work / ("bad-" + key + ".docx"), "--profile", path, fail=True)
    passed.append("plain documents are not declared format-checked; invalid profile values are rejected")
    print(json.dumps({"allPassed": True, "passed": passed}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
