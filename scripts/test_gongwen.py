"""Behavioral regression suite. Requires python-docx and Pillow for fixtures only."""
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import zipfile
from lxml import etree as ET

ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT / "scripts/gongwen.py"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"w": W, "m": M, "r": R}


def cmd(*args, failure=False):
    result = subprocess.run([sys.executable, str(CLI), *map(str, args)], capture_output=True,
                            text=True, encoding="utf-8", errors="replace")
    if failure:
        assert result.returncode, "Expected a safe refusal"
        return result.stderr
    assert result.returncode == 0, result.stdout + result.stderr
    return json.loads(result.stdout) if result.stdout.strip() else None


def contents(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def xml(files, name="word/document.xml"):
    return ET.fromstring(files[name])


def el(name, text=None, **attrs):
    node = ET.Element("{" + W + "}" + name)
    for k, v in attrs.items():
        node.set("{" + W + "}" + k, str(v))
    if text is not None:
        node.text = text
    return node


def fixture(path):
    from docx import Document
    from docx.shared import Inches
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from PIL import Image, ImageDraw
    image = path.parent / "fixture.png"
    im = Image.new("RGB", (960, 240), "#e6edf2")
    d = ImageDraw.Draw(im); d.rectangle((20, 20, 940, 220), outline="#355c70", width=4)
    d.text((50, 100), "IMAGE / retained media", fill="#243d4a")
    im.save(image)
    doc = Document()
    doc.add_heading("关于测试内容保留和跨页排版工作的通知", 0)
    doc.add_paragraph("各有关单位：")
    doc.add_heading("一、工作要求", 1)
    doc.add_paragraph("（一）总体要求。保留正文与题头字体的区别。")
    p = doc.add_paragraph("字符格式：")
    p.add_run("重点").bold = True
    p.add_run("下划线").underline = True
    p.add_run("2").font.superscript = True
    rel_id = doc.part.relate_to("https://example.com/reference", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "参考链接"; run.append(text); link.append(run); p._p.append(link)
    table = doc.add_table(rows=3, cols=3)
    table.autofit = False
    for column in table.columns:
        column.width = Inches(3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
    table.cell(0, 2).text = "值"
    for i in range(1, 3):
        for j in range(3):
            table.cell(i, j).text = f"单元格{i}-{j}，保持内容。"
    doc.add_picture(str(image), width=Inches(7))
    p = doc.add_paragraph("公式：")
    omath = ET.Element("{" + M + "}oMath")
    mr = ET.SubElement(omath, "{" + M + "}r"); ET.SubElement(mr, "{" + M + "}t").text = "x²+1"
    p._p.append(omath)
    doc.add_paragraph("自动编号事项甲", style="List Number")
    doc.add_paragraph("自动编号事项乙", style="List Number")
    for i in range(6):
        doc.add_paragraph(f"第{i+1}项说明。" + "这是用于检查跨页、段落缩进和内容保留的测试文字。" * 3)
    doc.add_paragraph("附件：1. 关于跨行附件名称保持对齐和完整内容的测试说明材料")
    doc.add_paragraph("测试工作办公室")
    doc.add_paragraph("2026年10月8日")
    doc.sections[0].header.paragraphs[0].text = "测试页眉（必须保留）"
    doc.sections[0].footer.paragraphs[0].text = "测试页脚（必须保留）"
    doc.save(path)
    files = contents(path)
    root = xml(files)
    p = root.find("w:body", NS).findall("w:p", NS)[4]
    p.append(el("commentRangeStart", id=0))
    crun = el("r"); crun.append(el("t", "被批注文字")); p.append(crun)
    p.append(el("commentRangeEnd", id=0))
    crun = el("r"); crun.append(el("commentReference", id=0)); p.append(crun)
    fnrun = el("r"); fnrun.append(el("footnoteReference", id=1)); p.append(fnrun)
    enrun = el("r"); enrun.append(el("endnoteReference", id=1)); p.append(enrun)
    ins = el("ins", id=42, author="测试", date="2026-10-08T00:00:00Z")
    ir = el("r"); ir.append(el("t", "新增文字")); ins.append(ir); p.append(ins)
    delete = el("del", id=43, author="测试", date="2026-10-08T00:00:00Z")
    dr = el("r"); dr.append(el("delText", "删除文字")); delete.append(dr); p.append(delete)
    files["word/document.xml"] = ET.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    comment_root = el("comments")
    comment = el("comment", id=0, author="测试", initials="T", date="2026-10-08T00:00:00Z")
    cp = el("p"); rr = el("r"); rr.append(el("t", "批注内容必须保留。")); cp.append(rr); comment.append(cp); comment_root.append(comment)
    new_parts = {"comments": comment_root}
    for plural, singular, text in (("footnotes", "footnote", "脚注原文"), ("endnotes", "endnote", "尾注原文")):
        notes = el(plural); note = el(singular, id=1); np = el("p"); nr = el("r"); nr.append(el("t", text)); np.append(nr); note.append(np); notes.append(note); new_parts[plural] = notes
    rels = xml(files, "word/_rels/document.xml.rels")
    types = xml(files, "[Content_Types].xml")
    for kind, part in new_parts.items():
        files[f"word/{kind}.xml"] = ET.tostring(part, xml_declaration=True, encoding="UTF-8", standalone=True)
        ET.SubElement(rels, "{" + PR + "}Relationship", Id="rIdTest" + kind, Type=R + "/" + kind, Target=kind + ".xml")
        ET.SubElement(types, "{" + CT + "}Override", PartName=f"/word/{kind}.xml", ContentType=f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}+xml")
    files["word/_rels/document.xml.rels"] = ET.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
    files["[Content_Types].xml"] = ET.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items(): z.writestr(name, data)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--work", required=True)
    args = parser.parse_args()
    work = Path(args.work).resolve()
    work.mkdir(parents=True, exist_ok=False)
    passed = []
    source = work / "rich.docx"; fixture(source)
    out = work / "rich-gbt.docx"
    cmd("format", source, "--output", out)
    checked = cmd("check", out, "--against", source)
    assert checked["contentPreserved"] and not checked["schemaErrors"]
    before, after = contents(source), contents(out)
    for part in ("word/comments.xml", "word/footnotes.xml", "word/endnotes.xml"):
        assert before[part] == after[part]
    b, a = xml(before), xml(after)
    formatted = a.xpath("//w:p[w:pPr/w:pStyle[starts-with(@w:val,'GW_')]]", namespaces=NS)
    assert formatted
    for paragraph in formatted:
        assert paragraph.xpath("w:pPr/w:spacing/@w:lineRule", namespaces=NS) == ["exact"]
        assert paragraph.xpath("w:pPr/w:spacing/@w:line", namespaces=NS) == ["578"]
        role = paragraph.xpath("w:pPr/w:pStyle/@w:val", namespaces=NS)[0]
        expected_snap = "0" if role in ("GW_title", "GW_attachment-title", "GW_signature", "GW_date") else "1"
        assert paragraph.xpath("w:pPr/w:snapToGrid/@w:val", namespaces=NS) == [expected_snap]
    for property_name, value in (("lineRule", "auto"), ("snapToGrid", "0")):
        corrupted = dict(after)
        bad_root = xml(corrupted)
        target = bad_root.xpath("//w:p[w:pPr/w:pStyle[@w:val='GW_body']]", namespaces=NS)[0]
        if property_name == "lineRule":
            target.find("w:pPr/w:spacing", NS).set("{" + W + "}lineRule", value)
        else:
            target.find("w:pPr/w:snapToGrid", NS).set("{" + W + "}val", value)
        corrupted["word/document.xml"] = ET.tostring(bad_root, xml_declaration=True, encoding="UTF-8", standalone=True)
        bad_file = work / ("bad-" + property_name + ".docx")
        with zipfile.ZipFile(bad_file, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in corrupted.items(): archive.writestr(name, data)
        cmd("check", bad_file, failure=True)
    passed.append("default spacing: fixed 28.9pt with body grid alignment; check rejects wrong line type or grid alignment")
    for expression in ("//m:oMath", "//w:ins", "//w:del", "//w:hyperlink", "//w:footnoteReference", "//w:endnoteReference", "//w:commentReference"):
        assert len(b.xpath(expression, namespaces=NS)) == len(a.xpath(expression, namespaces=NS)) > 0
    assert a.xpath("//w:vertAlign[@w:val='superscript']", namespaces=NS)
    assert a.xpath("//w:u", namespaces=NS)
    numbered = a.xpath("//w:p[w:r/w:t='自动编号事项甲']", namespaces=NS)[0]
    assert numbered.find("w:pPr/w:numPr/w:numId", NS) is not None
    assert sum(int(x.get("{" + W + "}w")) for x in a.xpath("//w:tblGrid/w:gridCol", namespaces=NS)) <= round(156 * 1440 / 25.4)
    assert all(after[n] == v for n, v in before.items() if n.startswith("word/media/"))
    passed.append("rich DOCX: media, table merges, formula, hyperlink, notes, comments, revisions and character emphasis preserved")

    second = work / "rich-twice.docx"
    cmd("format", out, "--output", second)
    cmd("check", second, "--against", out)
    again = contents(second)
    assert len([n for n in after if n.startswith("word/gongwen-footer-")]) == len([n for n in again if n.startswith("word/gongwen-footer-")]) == 3
    first_root, second_root = xml(after), xml(again)
    assert len(first_root.xpath("//w:p", namespaces=NS)) == len(second_root.xpath("//w:p", namespaces=NS))
    assert first_root.xpath("//w:rStyle[starts-with(@w:val,'GW_char_')]/@w:val", namespaces=NS) == second_root.xpath("//w:rStyle[starts-with(@w:val,'GW_char_')]/@w:val", namespaces=NS)
    passed.append("idempotence: paragraph count, footer count, inline heading styles and text stable")

    md = work / "input.md"
    md.write_text("# 关于做好年度工作安排并检查长标题换行与各级结构识别工作的通知\n\n各有关单位：\n\n## 工作原则\n\n（一）总体要求。保持原文编号。\n\n### 实施安排\n\n普通正文与 **强调内容**、[链接](https://example.com)。\n\n| 项目 | 内容 |\n|---|---|\n| 检查 | 保留表格 |\n\n![图示](fixture.png)\n\n公式：$x^2+1$\n\n脚注测试[^1]。\n\n[^1]: 脚注原文。\n\n测试办公室\n\n2026年10月8日\n", encoding="utf-8")
    prepared = work / "markdown-prepared.docx"; cmd("prepare", md, "--output", prepared)
    md_out = work / "markdown-gbt.docx"; cmd("format", prepared, "--output", md_out)
    cmd("check", md_out, "--against", prepared)
    root = xml(contents(md_out))
    h = root.xpath("//w:p[w:pPr/w:pStyle[@w:val='GW_heading1']]", namespaces=NS)
    assert any("".join(p.xpath(".//w:t/text()", namespaces=NS)) == "工作原则" for p in h)
    assert root.xpath("//m:oMath", namespaces=NS)
    passed.append("Markdown: unnumbered headings unchanged, image/table/link/footnote/math converted")

    inspected_md = cmd("inspect", prepared)
    title_para = inspected_md["paragraphs"][0]
    split_at = title_para["expectedText"].index("各级")
    title_map = work / "title-map.json"
    title_map.write_text(json.dumps({"version": 1, "sourceSha256": inspected_md["sourceSha256"], "paragraphs": [{"id": title_para["id"], "expectedText": title_para["expectedText"], "role": "title", "lineBreaks": [split_at]}]}, ensure_ascii=False), encoding="utf-8")
    title_out = work / "markdown-title-break.docx"
    cmd("format", prepared, "--output", title_out, "--map", title_map)
    cmd("check", title_out, "--against", prepared)
    tr = xml(contents(title_out))
    title_node = tr.xpath("//w:p[w:pPr/w:pStyle[@w:val='GW_title']]", namespaces=NS)[0]
    assert len(title_node.xpath(".//w:br", namespaces=NS)) == 1
    assert "".join(title_node.xpath(".//w:t/text()", namespaces=NS)) == title_para["expectedText"]
    passed.append("semantic title breaks: explicit word boundary, unchanged title text and valid Word structure")

    renumbered = work / "markdown-numbered.docx"
    cmd("format", prepared, "--output", renumbered, "--numbering", "official")
    texts = xml(contents(renumbered)).xpath("//w:t/text()", namespaces=NS)
    assert "一、工作原则" in texts and "（一）实施安排" in texts
    passed.append("explicit official numbering applies only to independent headings")

    txt = work / "input.txt"
    lines = ["关于文本排版的通知", "各单位：", "保留 * 星号、# 井号、1. 序号和 <标签>，不解释为Markdown。", "测试办公室", "2026年10月8日"]
    txt.write_text("\n".join(lines), encoding="utf-8")
    txt_out = work / "text-gbt.docx"; cmd("format", txt, "--output", txt_out)
    actual = ["".join(p.xpath(".//w:t/text()", namespaces=NS)) for p in xml(contents(txt_out)).xpath("/w:document/w:body/w:p", namespaces=NS)]
    assert actual == lines
    passed.append("TXT: literal Markdown-like characters and physical paragraphs preserved")

    calibration = work / "grid.txt"
    calibration.write_text("公文网格校准\n" + "天" * 1680, encoding="utf-8")
    cmd("format", calibration, "--output", work / "grid-gbt.docx")
    # Visual acceptance is performed by Word/PDF outside this structural suite.

    supplementary = work / "unicode.txt"
    supplementary.write_text("字符范围检查通知\n（一）𠮷字题头。后续正文。", encoding="utf-8")
    up = work / "unicode-prepared.docx"; cmd("prepare", supplementary, "--output", up)
    inspection = cmd("inspect", up)
    para = inspection["paragraphs"][1]
    lead = "（一）𠮷字题头。"
    mp = work / "unicode-map.json"
    mp.write_text(json.dumps({"version": 1, "sourceSha256": inspection["sourceSha256"], "paragraphs": [{"id": para["id"], "expectedText": para["expectedText"], "role": "body", "spans": [{"start": 0, "length": len(lead), "role": "heading2", "expectedText": lead}]}]}, ensure_ascii=False), encoding="utf-8")
    uo = work / "unicode-out.docx"; cmd("format", up, "--output", uo, "--map", mp)
    cmd("check", uo, "--against", up)
    ur = xml(contents(uo))
    assert "".join(ur.xpath("//w:r[w:rPr/w:rStyle[@w:val='GW_char_heading2']]/w:t/text()", namespaces=NS)) == lead
    passed.append("Unicode code-point mapping: supplementary CJK character does not shift inline heading boundaries")

    inspected = cmd("inspect", source)
    stale = work / "stale.json"; stale.write_text(json.dumps({"version": 1, "sourceSha256": "0" * 64, "paragraphs": []}), encoding="utf-8")
    refused = work / "refused.docx"
    cmd("format", source, "--output", refused, "--map", stale, failure=True); assert not refused.exists()
    digest = hashlib.sha256(out.read_bytes()).hexdigest()
    cmd("format", source, "--output", out, failure=True)
    assert hashlib.sha256(out.read_bytes()).hexdigest() == digest
    cmd("format", source, "--output", refused, "--font", "title=不存在的测试字体_abcdefgh", failure=True)
    assert not refused.exists()
    passed.append("safe failures: stale map, existing output and missing fonts do not publish or overwrite files")

    report = {"passed": passed, "count": len(passed), "allPassed": True}
    (work / "test-results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
