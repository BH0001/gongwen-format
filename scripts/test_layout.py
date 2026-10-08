"""Behavioral regression checks for footer content and section-aware sizing."""
import argparse
import json
from pathlib import Path
import subprocess
import sys
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Mm
from lxml import etree as ET
from PIL import Image

CLI = Path(__file__).resolve().with_name("gongwen.py")
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}


def page(paragraph, instruction=" PAGE "):
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "1"
    run.append(text); field.append(run); paragraph._p.append(field)


def complex_page(paragraph):
    for name, value in [("fldChar", "begin"), ("instrText", " PA"), ("instrText", "GE "),
                        ("fldChar", "separate"), ("t", "1"), ("fldChar", "end")]:
        run = OxmlElement("w:r"); node = OxmlElement("w:" + name)
        if name == "fldChar": node.set(qn("w:fldCharType"), value)
        else: node.text = value
        run.append(node); paragraph._p.append(run)


def contents(path):
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def xml(files, name="word/document.xml"):
    return ET.fromstring(files[name])


def text(node):
    return "".join(node.xpath(".//w:t/text()", namespaces=NS))


def base():
    doc = Document(); doc.add_paragraph("正文用于排版回归验证。")
    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--work", required=True)
    args = parser.parse_args()
    work = Path(args.work).resolve(); work.mkdir(parents=True, exist_ok=True)
    passed = []

    def format_source(source, target):
        result = subprocess.run([sys.executable, str(CLI), "format", str(source), "--output", str(target)],
                                capture_output=True, text=True, encoding="utf-8", errors="replace")
        assert result.returncode == 0, result.stdout + result.stderr
        report = json.loads(target.with_suffix(".report.json").read_text(encoding="utf-8"))
        assert report["contentPreserved"] and not report["newSchemaErrors"] and not report["formatErrors"]
        return contents(target)

    def format_doc(name, doc):
        source = work / (name + ".docx"); target = work / (name + "-out.docx")
        assert not source.exists(), "Use an empty work directory"
        doc.save(source)
        return format_source(source, target)

    def footer(files):
        return xml(files, "word/gongwen-footer-1-default.xml")

    def rejects_changed_content(name, files, source):
        target = work / (name + "-tampered.docx")
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
            for part, data in files.items(): archive.writestr(part, data)
        result = subprocess.run([sys.executable, str(CLI), "check", str(target), "--against", str(source)],
                                capture_output=True, text=True, encoding="utf-8", errors="replace")
        assert result.returncode != 0 and '"contentPreserved": false' in result.stdout + result.stderr, result.stdout + result.stderr

    doc = base(); p = doc.sections[0].footer.paragraphs[0]
    p.text = "2026/10/08 — "; page(p); p.add_run(" —")
    assert text(footer(format_doc("numeric-footer", doc))) == "2026/10/08 — 1 —"
    passed.append("Date and identifier digits outside PAGE are retained")

    doc = base(); table = doc.sections[0].footer.add_table(rows=1, cols=1, width=Mm(80))
    page(table.cell(0, 0).paragraphs[0]); node = footer(format_doc("table-footer", doc))
    assert len(node.xpath(".//w:fldSimple | .//w:instrText", namespaces=NS)) == 1
    passed.append("PAGE inside footer tables does not create a duplicate")

    doc = base(); page(doc.sections[0].footer.paragraphs[0], " page ")
    assert text(footer(format_doc("lowercase-page", doc))) == "—1—"
    passed.append("Case-insensitive PAGE field recognition")

    doc = base(); complex_page(doc.sections[0].footer.paragraphs[0])
    assert text(footer(format_doc("split-page", doc))) == "—1—"
    passed.append("PAGE instruction split across runs is recognized")

    doc = base(); p = doc.sections[0].footer.paragraphs[0]; p.text = "编号 20261008 "; complex_page(p)
    assert text(footer(format_doc("mixed-page", doc))) == "编号 20261008 1"
    passed.append("Content beside complex PAGE fields is retained")

    doc = base(); p = doc.sections[0].footer.paragraphs[0]; page(p); p.add_run("/"); page(p, " NUMPAGES ")
    assert footer(format_doc("page-count", doc)).xpath("count(.//w:fldSimple)", namespaces=NS) == 2
    passed.append("PAGE/NUMPAGES combinations are retained without additions")

    doc = base(); sec = doc.sections[0]
    sec.header.paragraphs[0].text = "ACTIVE HEADER"; sec.footer.paragraphs[0].text = "ACTIVE FOOTER"
    sec.even_page_header.paragraphs[0].text = "STALE HEADER"; sec.even_page_footer.paragraphs[0].text = "STALE FOOTER"
    setting = OxmlElement("w:evenAndOddHeaders"); setting.set(qn("w:val"), "false"); doc.settings.element.append(setting)
    files = format_doc("inactive-even", doc); root = xml(files); rels = xml(files, "word/_rels/document.xml.rels")
    rid = root.xpath("//w:headerReference[@w:type='even']/@r:id", namespaces=NS)[0]
    target = next(node.get("Target") for node in rels if node.get("Id") == rid)
    assert text(xml(files, "word/" + target)) == "ACTIVE HEADER"
    even_text = text(xml(files, "word/gongwen-footer-1-even.xml"))
    assert "ACTIVE FOOTER" in even_text and "STALE" not in even_text
    passed.append("Enabling alternating numbers does not activate inactive even-page text")

    wide = work / "wide.png"; tall = work / "tall.png"
    Image.new("RGB", (1000, 100), "red").save(wide); Image.new("RGB", (100, 1000), "red").save(tall)
    doc = base(); sec = doc.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Mm(297); sec.page_height = Mm(210)
    table = doc.add_table(rows=1, cols=2); table.autofit = False
    for col in table.columns: col.width = Mm(100)
    doc.add_picture(str(wide), width=Mm(200)); doc.add_picture(str(tall), height=Mm(200))
    root = xml(format_doc("landscape", doc))
    assert [int(v) for v in root.xpath("//w:tblGrid/w:gridCol/@w:w", namespaces=NS)] == [5669, 5669]
    assert [int(v) for v in root.xpath("//wp:extent/@cx", namespaces=NS)][0] == 7200000
    assert [int(v) for v in root.xpath("//wp:extent/@cy", namespaces=NS)][1] == 4968240
    passed.append("Landscape tables and pictures use landscape width and height")

    doc = base(); doc.add_picture(str(wide), width=Mm(200)); sec = doc.add_section()
    sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width = Mm(297); sec.page_height = Mm(210)
    doc.add_picture(str(wide), width=Mm(200)); root = xml(format_doc("mixed-sections", doc))
    assert [int(v) for v in root.xpath("//wp:extent/@cx", namespaces=NS)] == [5615940, 7200000]
    passed.append("Mixed orientation documents use each object's containing section")

    doc = base(); doc.sections[0].footer.paragraphs[0].add_run().add_picture(str(wide), width=Mm(40))
    source = work / "nested-footer.docx"; doc.save(source); files = contents(source)
    files["word/nested/footer1.xml"] = files.pop("word/footer1.xml")
    rels = xml(files, "word/_rels/footer1.xml.rels")
    for node in rels: node.set("Target", "../" + node.get("Target"))
    files["word/nested/_rels/footer1.xml.rels"] = ET.tostring(rels)
    del files["word/_rels/footer1.xml.rels"]
    rels = xml(files, "word/_rels/document.xml.rels")
    for node in rels:
        if node.get("Target") == "footer1.xml": node.set("Target", "nested/footer1.xml")
    files["word/_rels/document.xml.rels"] = ET.tostring(rels)
    types = xml(files, "[Content_Types].xml")
    for node in types:
        if node.get("PartName") == "/word/footer1.xml": node.set("PartName", "/word/nested/footer1.xml")
    files["[Content_Types].xml"] = ET.tostring(types)
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items(): archive.writestr(name, data)
    target = work / "nested-footer-out.docx"; produced = format_source(source, target)
    rels = xml(produced, "word/_rels/gongwen-footer-1-default.xml.rels")
    assert any(node.get("Target", "").startswith("media/") for node in rels)
    assert footer(produced).xpath("count(.//w:drawing)", namespaces=NS) == 1
    format_source(target, work / "nested-footer-twice.docx")
    passed.append("Footer relationships in nested directories survive relocation and a second format")

    # The original source footer remains in the package. Damage only its active
    # replacement to prove that checking orphan source parts cannot mask loss.
    files = contents(work / "numeric-footer-out.docx"); node = footer(files)
    node.xpath(".//w:t", namespaces=NS)[0].text = ""
    files["word/gongwen-footer-1-default.xml"] = ET.tostring(node)
    rejects_changed_content("numeric-footer", files, work / "numeric-footer.docx")
    passed.append("Content check rejects removal from the active copied footer")

    doc = base(); footer_part = doc.sections[0].footer.part
    rid = footer_part.relate_to("https://example.com/original", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), rid)
    run = OxmlElement("w:r"); label = OxmlElement("w:t"); label.text = "链接内容"
    run.append(label); link.append(run); doc.sections[0].footer.paragraphs[0]._p.append(link)
    files = format_doc("linked-footer", doc)
    rels = xml(files, "word/_rels/gongwen-footer-1-default.xml.rels")
    for node in rels:
        if node.get("Type", "").endswith("/hyperlink"): node.set("Target", "https://example.com/changed")
    files["word/_rels/gongwen-footer-1-default.xml.rels"] = ET.tostring(rels)
    rejects_changed_content("linked-footer", files, work / "linked-footer.docx")
    passed.append("Content check resolves active footer relationships and rejects changed hyperlink targets")

    doc = base(); sec = doc.sections[0]; sec.different_first_page_header_footer = True
    sec.header.paragraphs[0].text = "DEFAULT HEADER"; sec.first_page_header.paragraphs[0].text = "FIRST HEADER"
    sec.footer.paragraphs[0].text = "DEFAULT FOOTER"; sec.first_page_footer.paragraphs[0].text = "FIRST FOOTER"
    doc.add_section(); doc.add_paragraph("第二节沿用页眉页脚。")
    files = format_doc("inherited-first-footer", doc)
    node = xml(files, "word/gongwen-footer-2-first.xml")
    node.xpath(".//w:t", namespaces=NS)[0].text = ""
    files["word/gongwen-footer-2-first.xml"] = ET.tostring(node)
    rejects_changed_content("inherited-first-footer", files, work / "inherited-first-footer.docx")
    passed.append("Section inheritance and first-page activation are checked for content retention")

    result = {"passed": passed, "count": len(passed), "allPassed": True}
    (work / "test-results.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
